"""Minimal slide-deck builder — PowerPoint under the hood, PDF as the deliverable.

Shared by the decks in scripts/, so layout and styling live in one place.

PDF export goes through Keynote via AppleScript. LibreOffice would be the usual
headless route, but the `soffice` on this machine is a broken wrapper pointing at
an application that is not installed.
"""

import subprocess
import time
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

INK = RGBColor(0x0B, 0x0B, 0x0B)
INK2 = RGBColor(0x52, 0x51, 0x4E)
MUTED = RGBColor(0x8A, 0x8A, 0x85)
BLUE = RGBColor(0x2A, 0x78, 0xD6)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
AQUA = RGBColor(0x1B, 0xAF, 0x7A)
RED = RGBColor(0xE3, 0x49, 0x48)
BG = RGBColor(0xFC, 0xFC, 0xFB)
LINE = RGBColor(0xE3, 0xE3, 0xDF)
HEAD_BG = RGBColor(0xF2, 0xF5, 0xFA)
FONT = "Helvetica Neue"

W, H = 13.333, 7.5


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self._blank = self.prs.slide_layouts[6]
        self.n = 0

    # ---------------- slides ----------------

    def slide(self, kicker=None, title=None, sub=None):
        s = self.prs.slides.add_slide(self._blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = BG
        self.n += 1
        if title:
            self.text(s, .7, .45, 12, .3, kicker.upper(), size=12, color=BLUE, bold=True)
            # A title that wraps to two lines used to run straight through the rule
            # and the subtitle, so shrink it until it fits on one.
            size = 34 if len(title) <= 46 else (29 if len(title) <= 58 else 25)
            self.text(s, .7, .82, 12.1, .7, title, size=size, bold=True)
            self.rule(s, .7, 1.58, 2.2, BLUE, .04)
            if sub:
                self.text(s, .7, 1.78, 12.1, .5, sub, size=15, color=INK2)
        if self.n > 1:
            self.text(s, 11.9, 6.95, 1.0, .3, str(self.n), size=11, color=MUTED,
                      align=PP_ALIGN.RIGHT)
        return s

    def title_slide(self, kicker, title, subtitle, author, note=None):
        s = self.slide()
        self.rule(s, 0, 0, W, BLUE, .1)
        self.text(s, 1.0, 2.5, 11.4, .35, kicker, size=14, color=BLUE, bold=True)
        self.text(s, 1.0, 3.0, 11.6, 1.2, title, size=48, bold=True)
        self.rule(s, 1.0, 4.45, 2.6, BLUE, .045)
        self.text(s, 1.0, 4.75, 11.4, .8, subtitle, size=17, color=INK2)
        self.text(s, 1.0, 6.5, 8, .35, author, size=14, bold=True)
        if note:
            self.text(s, 1.0, 6.82, 10, .3, note, size=11, color=MUTED)
        return s

    # ---------------- elements ----------------

    def text(self, s, x, y, w, h, body, size=18, color=INK, bold=False,
             align=PP_ALIGN.LEFT, space_after=10, line=1.25):
        box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, ln in enumerate(body.split("\n") if isinstance(body, str) else body):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment, p.space_after, p.line_spacing = align, Pt(space_after), line
            r = p.add_run()
            r.text = ln
            r.font.size, r.font.bold = Pt(size), bold
            r.font.color.rgb, r.font.name = color, FONT
        return box

    def rule(self, s, x, y, w, color=LINE, h=0.03):
        sh = s.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = color
        sh.line.fill.background()
        sh.shadow.inherit = False

    def image(self, s, path, x, y, w, h):
        """Scale to fit inside the box, centred, aspect preserved."""
        p = Path(path)
        if not p.exists():
            self.text(s, x, y, 4, .4, f"[missing {p.name}]", size=12, color=RED)
            return
        iw, ih = Image.open(p).size
        sc = min(w / iw, h / ih)
        ww, hh = iw * sc, ih * sc
        s.shapes.add_picture(str(p), Inches(x + (w - ww) / 2), Inches(y + (h - hh) / 2),
                             width=Inches(ww), height=Inches(hh))

    def table(self, s, x, y, w, rows, col_w=None, size=14, row_h=.46,
              highlight=()):
        """Draw a table; returns the y coordinate of its bottom edge.

        Use that return value to place whatever follows — hardcoding a y below a
        table is how text ends up sitting on the last row.

        highlight: row indices (1-based over data rows) to emphasise.
        """
        nr, nc = len(rows), len(rows[0])
        t = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w),
                               Inches(row_h * nr)).table
        if col_w:
            tot = sum(col_w)
            for i, cw in enumerate(col_w):
                t.columns[i].width = Emu(int(Inches(w) * cw / tot))
        for i, row in enumerate(rows):
            t.rows[i].height = Inches(row_h)
            for j, val in enumerate(row):
                c = t.cell(i, j)
                c.text = str(val)
                c.margin_left = c.margin_right = Inches(.12)
                c.vertical_anchor = MSO_ANCHOR.MIDDLE
                c.fill.solid()
                if i == 0:
                    c.fill.fore_color.rgb = HEAD_BG
                elif i in highlight:
                    c.fill.fore_color.rgb = RGBColor(0xFD, 0xEE, 0xE6)
                else:
                    c.fill.fore_color.rgb = BG
                # style every paragraph, not just the first: a cell containing a
                # newline gets a second paragraph, which otherwise keeps the
                # default theme font and renders in the wrong typeface
                for para in c.text_frame.paragraphs:
                    for r in para.runs:
                        r.font.size = Pt(size)
                        r.font.bold = (i == 0) or (i in highlight)
                        r.font.color.rgb = INK if (i == 0 or i in highlight) else INK2
                        r.font.name = FONT
        return y + row_h * nr

    # ---------------- output ----------------

    def save_pdf(self, pdf_path, keep_pptx=False):
        pdf = Path(pdf_path)
        pdf.parent.mkdir(parents=True, exist_ok=True)
        pptx = pdf.with_suffix(".pptx")
        self.prs.save(str(pptx))
        if pdf.exists():
            pdf.unlink()
        # Keynote occasionally returns "User cancelled" if it is still shutting
        # down from a previous export, so try again once before giving up.
        script = f'''
tell application "Keynote"
    launch
    set d to open (POSIX file "{pptx}")
    delay 3
    export d to (POSIX file "{pdf}") as PDF with properties {{PDF image quality:Best}}
    close d saving no
    quit
end tell'''
        for attempt in (1, 2):
            r = subprocess.run(["osascript", "-e", script],
                               capture_output=True, text=True)
            if r.returncode == 0 and pdf.exists():
                break
            if attempt == 1:
                subprocess.run(["pkill", "-f", "Keynote"], capture_output=True)
                time.sleep(3)
            else:
                raise RuntimeError(f"Keynote export failed: {r.stderr.strip()}")
        if not keep_pptx:
            pptx.unlink()
        print(f"wrote {pdf}  ({pdf.stat().st_size / 1e6:.1f} MB, {self.n} slides)")
        return pdf
